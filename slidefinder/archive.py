"""Phase 2: read raw decks from the lecture archive.

    python -m slidefinder.archive 45-851            # one course, every term
    python -m slidefinder.archive 45-851 --limit 3  # first three sessions, for timing

For each session folder `<NN> <date> <title>` under `<course>/<term>/`, this renders every page of
slides.pdf (or slides.pptx converted with LibreOffice) to WebP, pulls the page text and the pptx
speaker notes, and writes one row per page to build/<course>.jsonl plus the images under
build/slides/<course>/. Nothing here talks to the database; see dedup.py and load_archive.py.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

import fitz  # PyMuPDF
from PIL import Image

from . import config

BUILD = config.ROOT / "build"
IMAGE_WIDTH, THUMB_WIDTH = 1600, 400
SESSION_RE = re.compile(r"^(\d+) (\d{4}-\d{2}-\d{2}) (.+)$")
SEASON = {"Spring": 1, "Summer": 2, "Fall": 3}

ARCHIVE_COURSES = {
    "45-851": "45-851 Data Mining",
    "45-885": "45-885 Data Visualization",
    "46-887": "46-887 Machine Learning for Business Applications",
    "45-884": "45-884 AI Methods for Social and Visual Data",
    "70-445": "70-445 AI for Business Leaders",
}
# Sessions whose decks are mostly students' own work; their slides stay out of the index.
SKIP_TITLE = re.compile(r"student present|project present|final present|team present", re.I)


def term_key(term: str) -> tuple[int, int, str]:
    """'2025 Spring (M4)' -> (2025, 1, 'M4'), so terms sort chronologically."""
    m = re.match(r"(\d{4}) (Spring|Summer|Fall)\s*(?:\((.*?)\))?", term)
    return (int(m.group(1)), SEASON[m.group(2)], m.group(3) or "") if m else (0, 0, term)


def term_slug(term: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", term.lower())


def soffice() -> str | None:
    for c in (shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              str(Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice")):
        if c and Path(c).exists():
            return c
    return None


def pptx_to_pdf(pptx: Path, out_dir: Path) -> Path | None:
    exe = soffice()
    if not exe:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / (pptx.stem + ".pdf")
    if target.exists() and target.stat().st_mtime > pptx.stat().st_mtime:
        return target
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx)],
                   capture_output=True, timeout=600)
    return target if target.exists() else None


def pptx_notes(pptx: Path) -> list[str]:
    try:
        from pptx import Presentation
        prs = Presentation(str(pptx))
        out = []
        for s in prs.slides:
            if s.element.get("show") == "0":  # hidden slide, not in the PDF
                continue
            tf = s.notes_slide.notes_text_frame if s.has_notes_slide else None
            out.append(" ".join((tf.text if tf else "").split()))
        return out
    except Exception as e:  # a broken pptx should not stop the run
        print(f"  notes failed for {pptx.name}: {e}", file=sys.stderr)
        return []


def dhash(im: Image.Image) -> int:
    g = im.convert("L").resize((9, 8), Image.Resampling.LANCZOS)
    px = list(g.getdata())  # noqa
    bits = 0
    for r in range(8):
        for c in range(8):
            bits = (bits << 1) | (px[r * 9 + c] > px[r * 9 + c + 1])
    return bits


def clean_title(first_line: str) -> str:
    t = " ".join(first_line.split())[:140]
    return t if re.search(r"[A-Za-z]{3,}", t) else ""  # OCR noise like "a0D OD" is not a title


def iter_sessions(course: str):
    root = config.LECTURE_ARCHIVE / ARCHIVE_COURSES[course]
    for term_dir in sorted(root.iterdir()):
        if not term_dir.is_dir() or not re.match(r"\d{4} ", term_dir.name):
            continue
        for sess in sorted(term_dir.iterdir()):
            m = SESSION_RE.match(sess.name)
            if not sess.is_dir() or not m:
                continue
            pdf, pptx = sess / "slides.pdf", sess / "slides.pptx"
            if not pdf.exists() and not pptx.exists():
                continue
            yield {"term": term_dir.name, "session": int(m.group(1)), "date": m.group(2),
                   "title": m.group(3), "dir": sess, "pdf": pdf if pdf.exists() else None,
                   "pptx": pptx if pptx.exists() else None}


def render_session(course: str, s: dict, out_dir: Path) -> list[dict]:
    pdf = s["pdf"] or pptx_to_pdf(s["pptx"], BUILD / "converted" / course / term_slug(s["term"]))
    if not pdf:
        print(f"  skip (pptx only, no LibreOffice): {s['dir'].name}", file=sys.stderr)
        return []
    doc = fitz.open(pdf)
    notes = pptx_notes(s["pptx"]) if s["pptx"] else []
    if notes and len(notes) != len(doc):
        notes = []  # counts differ, so index-matching would attach the wrong notes
    rows = []
    prefix = f"{course.replace('-', '')}-{term_slug(s['term'])}-s{s['session']:02d}"
    for i, page in enumerate(doc):
        n = i + 1
        sid = f"{prefix}-{n:03d}"
        text = page.get_text().strip()
        lines = [l for l in text.splitlines() if l.strip()]
        pix = page.get_pixmap(matrix=fitz.Matrix(IMAGE_WIDTH / page.rect.width, IMAGE_WIDTH / page.rect.width), alpha=False)
        im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img_path, thumb_path = out_dir / f"{sid}.webp", out_dir / f"{sid}-thumb.webp"
        if not (img_path.exists() and thumb_path.exists()):
            im.save(img_path, "WEBP", quality=80)
            im.resize((THUMB_WIDTH, round(im.height * THUMB_WIDTH / im.width)), Image.Resampling.LANCZOS).save(thumb_path, "WEBP", quality=75)
        rows.append({
            "slide_id": sid, "course": course, "term": s["term"], "session": s["session"],
            "session_title": s["title"], "session_date": s["date"], "slide_number": n,
            "title": clean_title(lines[0]) if lines else "", "text": text,
            "notes": notes[i] if i < len(notes) else "",
            "dhash": dhash(im), "text_hash": hashlib.sha1(" ".join(text.lower().split()).encode()).hexdigest()[:16],
            "image_path": f"slides/{course.replace('-', '')}/{sid}.webp",
            "thumb_path": f"slides/{course.replace('-', '')}/{sid}-thumb.webp",
            "local_image": str(img_path), "local_thumb": str(thumb_path),
        })
    return rows


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("course", choices=ARCHIVE_COURSES)
    p.add_argument("--limit", type=int, help="only the first N sessions")
    args = p.parse_args()

    out_dir = BUILD / "slides" / args.course.replace("-", "")
    out_dir.mkdir(parents=True, exist_ok=True)
    sessions = list(iter_sessions(args.course))
    sessions = [s for s in sessions if not SKIP_TITLE.search(s["title"])]
    if args.limit:
        sessions = sessions[:args.limit]
    print(f"{args.course}: {len(sessions)} sessions across {len({s['term'] for s in sessions})} terms")
    t0 = time.time()
    rows: list[dict] = []
    for k, s in enumerate(sessions, 1):
        t = time.time()
        r = render_session(args.course, s, out_dir)
        rows.extend(r)
        print(f"  [{k}/{len(sessions)}] {s['term']} s{s['session']:02d} {s['title'][:45]}: {len(r)} pages, {time.time() - t:.1f}s", file=sys.stderr)
    out = BUILD / f"{args.course}.jsonl"
    with out.open("w") as f:
        for r in rows:
            f.write(json.dumps(r) + "\n")
    print(f"wrote {len(rows)} slides to {out} in {time.time() - t0:.0f}s")


if __name__ == "__main__":
    main()
